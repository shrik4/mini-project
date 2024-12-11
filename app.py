import streamlit as st 
import pyttsx3  # for Text-to-Speech Feature
from fpdf import FPDF  # for Export to PDF Feature
from PyPDF2 import PdfReader
from docx import Document
import pandas as pd
from langchain.text_splitter import RecursiveCharacterTextSplitter
import os
import sqlite3
from datetime import datetime, timedelta
from langchain_google_genai import GoogleGenerativeAIEmbeddings
import google.generativeai as genai
from langchain_community.vectorstores import FAISS
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.chains.question_answering import load_qa_chain
from langchain.prompts import PromptTemplate
from dotenv import load_dotenv

st.set_page_config(page_title="Chat PDF")

load_dotenv()
os.getenv("GOOGLE_API_KEY")
genai.configure(api_key=os.getenv("GOOGLE_API_KEY"))

# Initialize the SQLite database with a table for chat history
def init_db():
    conn = sqlite3.connect("chat_history.db")
    c = conn.cursor()
    c.execute('''CREATE TABLE IF NOT EXISTS chat_history
                 (id INTEGER PRIMARY KEY, question TEXT, answer TEXT, timestamp TEXT)''')
    conn.commit()
    conn.close()

# Save a question and answer to the database with the current timestamp
def save_to_db(question, answer):
    conn = sqlite3.connect("chat_history.db")
    c = conn.cursor()
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    c.execute("INSERT INTO chat_history (question, answer, timestamp) VALUES (?, ?, ?)", 
              (question, answer, timestamp))
    conn.commit()
    conn.close()

# Retrieve the chat history from the database and categorize by time
def get_chat_history():
    conn = sqlite3.connect("chat_history.db")
    c = conn.cursor()
    c.execute("SELECT question, answer, timestamp FROM chat_history ORDER BY id DESC")
    history = c.fetchall()
    conn.close()

    today = datetime.now().date()
    yesterday = today - timedelta(days=1)
    last_7_days = today - timedelta(days=7)
    last_30_days = today - timedelta(days=30)

    categorized_history = {
        "Today": [],
        "Yesterday": [],
        "Previous 7 Days": [],
        "Previous 30 Days": []
    }

    for question, answer, timestamp in history:
        ts = datetime.strptime(timestamp, "%Y-%m-%d %H:%M:%S").date()
        if ts == today:
            categorized_history["Today"].append((question, answer, timestamp))
        elif ts == yesterday:
            categorized_history["Yesterday"].append((question, answer, timestamp))
        elif ts >= last_7_days:
            categorized_history["Previous 7 Days"].append((question, answer, timestamp))
        elif ts >= last_30_days:
            categorized_history["Previous 30 Days"].append((question, answer, timestamp))

    return categorized_history

# Display categorized chat history
def display_chat_history():
    st.sidebar.subheader("Chat History")
    categorized_history = get_chat_history()

    for category, chats in categorized_history.items():
        if chats:
            st.sidebar.subheader(category)
            for question, answer, timestamp in chats:
                st.sidebar.write(f"**Question:** {question}")
                st.sidebar.write(f"**Answer:** {answer}")
                st.sidebar.write(f"**Timestamp:** {timestamp}")
                st.sidebar.write("---")

# Other functions remain unchanged...
def get_pdf_text(pdf_docs):
    text = ""
    for pdf in pdf_docs:
        pdf_reader = PdfReader(pdf)
        for page in pdf_reader.pages:
            text += page.extract_text()
    return text

def get_docx_text(docx_docs):
    text = ""
    for docx in docx_docs:
        doc = Document(docx)
        for para in doc.paragraphs:
            text += para.text + "\n"
    return text

def get_excel_text(excel_docs):
    text = ""
    for excel in excel_docs:
        df = pd.read_excel(excel)
        text += df.to_string(index=False)
    return text

def get_text_chunks(text):
    text_splitter = RecursiveCharacterTextSplitter(chunk_size=10000, chunk_overlap=1000)
    chunks = text_splitter.split_text(text)
    return chunks

def get_vector_store(text_chunks):
    embeddings = GoogleGenerativeAIEmbeddings(model="models/embedding-001")
    if not text_chunks:
        st.error("No text chunks were generated from the files.")
        return None
    vector_store = FAISS.from_texts(text_chunks, embedding=embeddings)
    vector_store.save_local("faiss_index")

def get_conversational_chain():
    prompt_template = """
    Answer the question as detailed as possible from the provided context. If the answer is not in
    provided context, just say "answer is not available in the context". Do not provide incorrect answers.\n\n
    Context:\n{context}\n
    Question:\n{question}\n
    Answer:
    """
    model = ChatGoogleGenerativeAI(model="gemini-pro", temperature=0.3)
    prompt = PromptTemplate(template=prompt_template, input_variables=["context", "question"])
    chain = load_qa_chain(model, chain_type="stuff", prompt=prompt)
    return chain

def text_to_speech(text):
    engine = pyttsx3.init()
    engine.say(text)
    engine.runAndWait()

def clean_text_for_pdf(text):
    return text.encode("latin-1", errors="replace").decode("latin-1")

def export_to_pdf(user_question, response):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Chat with Files using Gemini", ln=True, align='C')
    
    user_question_cleaned = clean_text_for_pdf(user_question)
    output_text_cleaned = clean_text_for_pdf(response["output_text"])

    pdf.cell(200, 10, txt="User Question: " + user_question_cleaned, ln=True)
    pdf.multi_cell(0, 10, txt="AI Answer: " + output_text_cleaned)
    pdf.output("chat_history.pdf")

def user_input(user_question):
    embeddings = GoogleGenerativeAIEmbeddings(model="models/embedding-001")
    new_db = FAISS.load_local("faiss_index", embeddings, allow_dangerous_deserialization=True)
    docs = new_db.similarity_search(user_question)
    chain = get_conversational_chain()
    response = chain({"input_documents": docs, "question": user_question}, return_only_outputs=True)
    st.write("Reply: ", response["output_text"])

    save_to_db(user_question, response["output_text"])

    if st.button("Read Answer Aloud"):
        text_to_speech(response["output_text"])

    if st.button("Export Chat to PDF"):
        export_to_pdf(user_question, response)

def main():
    init_db()
    st.header("Smart pdf Chat using Gemini🤖")

    user_question = st.text_input("Ask a Question from the Files")

    if user_question:
        user_input(user_question)

    st.markdown(
        """
        <p>Also follow us on <a href="https://www.linkedin.com/in/shrikar-shettigar-2971512a6/" target="_blank">LinkedIn</a>🚀 and <a href="https://github.com/shrik4" target="_blank">GitHub</a>📛</p>
        """,
        unsafe_allow_html=True
    )

    with st.sidebar:
        st.title("Menu:")
        uploaded_files = st.file_uploader("Upload your PDF, DOCX, or Excel Files and Click on the Submit & Process Button", accept_multiple_files=True, type=['pdf', 'docx', 'xlsx'])
        if st.button("Submit & Process"):
            with st.spinner("Processing..."):
                raw_text = ""
                for file in uploaded_files:
                    if file.type == "application/pdf":
                        raw_text += get_pdf_text([file])
                    elif file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
                        raw_text += get_docx_text([file])
                    elif file.type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
                        raw_text += get_excel_text([file])
                
                text_chunks = get_text_chunks(raw_text)
                if text_chunks:
                    get_vector_store(text_chunks)
                    st.success("Done")
                else:
                    st.error("No text chunks to process.")

    display_chat_history()

if __name__ == "__main__":
    main()
